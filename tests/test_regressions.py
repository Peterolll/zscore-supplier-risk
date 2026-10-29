import math
import unittest
from unittest.mock import patch
from zscore_pipeline.models import SupplierProfile, ExtractionResult, FieldValue, ExtractMethod, Gaap, PeriodType
from zscore_pipeline.pipeline import process
from zscore_pipeline.m6_calc import calc
from zscore_pipeline.m3_locate import period_factor
from zscore_pipeline.m4_map import resolve
from zscore_pipeline.m2_extract import LineItem, table_items, attach_page_metadata, finalize_units, _parse_ocr_text
from zscore_pipeline import m2_extract as m2


class Regressions(unittest.TestCase):
    def profile(self, **kwargs):
        return SupplierProfile(name='audit', industry_class='manufacturing', period_type='annual', gaap='cas', **kwargs)

    def test_production_never_reads_benchmark_overrides(self):
        p = self.profile().model_copy(update={'name': 'QSP 2022'})
        with patch('zscore_pipeline.pipeline.get_override', side_effect=AssertionError('benchmark accessed')):
            with self.assertRaises(FileNotFoundError):
                process(p, '/does-not-exist-audit.pdf')

    def test_every_required_factor_is_required(self):
        values = dict(current_assets=600, current_liabilities=200, total_assets=1000,
                      total_liabilities=400, equity_total=600, retained_earnings=200, revenue=1000, ebit=100)
        for key in values:
            fields = {k: FieldValue(value=v) for k, v in values.items() if k != key}
            ext = ExtractionResult(profile_name='audit', method='pdfplumber', fields=fields)
            score = calc(self.profile(), ext, 1)
            self.assertIsNone(score.z_score, key)
            self.assertEqual(score.risk_zone, 'unknown')
        ext = ExtractionResult(profile_name='audit', method='pdfplumber', fields={k: FieldValue(value=v) for k,v in values.items()})
        self.assertAlmostEqual(calc(self.profile(), ext, 1).z_score, 2.3949)
        ext.fields.pop('revenue')
        service=self.profile().model_copy(update={'industry_class':'service'})
        self.assertIsNotNone(calc(service,ext,1).z_score)

    def test_nonfinite_is_unscorable(self):
        ext=ExtractionResult(profile_name='audit',method='pdfplumber',fields={'total_assets':FieldValue(value=1),'revenue':FieldValue(value=math.inf)})
        self.assertIsNone(calc(self.profile(),ext,1).z_score)

    def test_table_preserves_missing_current_and_skips_note_column(self):
        rows=[['项目','附注','期末余额','期初余额'],['资产总计','5',None,'900000'],['营业收入','6','400000','300000']]
        out=table_items(rows,0,ExtractMethod.PDFPLUMBER)
        self.assertEqual([(x.label,x.raw_value) for x in out],[('资产总计',''),('营业收入','400000')])
        out=table_items([['资产总计',None,'900000']],0,ExtractMethod.PDFPLUMBER)
        self.assertEqual(out[0].raw_value,'')
        out=table_items([['项目','2024','2025'],['资产总计','123','456']],0,ExtractMethod.PDFPLUMBER)
        self.assertEqual(out[0].raw_value,'456')

    def test_real_small_amount_survives(self):
        items=[LineItem('资产总计','1000000',0,ExtractMethod.PDFPLUMBER),LineItem('利息费用','500',0,ExtractMethod.PDFPLUMBER)]
        result=resolve(items,set(),Gaap.CAS,1,'audit',ExtractMethod.PDFPLUMBER)
        self.assertEqual(result.fields['interest_expense'].value,500)

    def test_last_ocr_attempt_is_merged(self):
        with patch.dict(m2.GLM_OCR,{'verify_attempts':2,'verify_fields':['revenue','profit_before_tax']}), patch.object(m2,'ocr_pdf_via_glm',side_effect=['单位：元\n营业收入\t1000000','单位：元\n利润总额\t200000']):
            items,_,count,missing=m2._ocr_with_verification('unused')
        self.assertEqual(count,2)
        self.assertEqual(missing,[])
        self.assertEqual(len(items),2)

    def test_page_units_and_unknown_unit(self):
        a=LineItem('资产总计','100',0,ExtractMethod.PDFPLUMBER)
        b=LineItem('营业收入','2',1,ExtractMethod.PDFPLUMBER)
        c=LineItem('利息费用','100',2,ExtractMethod.PDFPLUMBER)
        attach_page_metadata([a],'资产负债表\n单位：万元')
        attach_page_metadata([b],'Income statement\nin millions')
        finalize_units([a,b,c])
        ext=resolve([a,b,c],set(),Gaap.CAS,1,'audit',ExtractMethod.PDFPLUMBER)
        self.assertEqual(ext.fields['total_assets'].value,1000000)
        self.assertEqual(ext.fields['revenue'].value,2000000)
        self.assertNotIn('interest_expense',ext.fields)
        self.assertTrue(any('单位未确认' in w for w in ext.warnings))

    def test_ocr_units_keep_real_page(self):
        items=_parse_ocr_text('[PAGE 3]\n单位：千元\n资产总计\t100\n[PAGE 5]\n单位：万元\n营业收入\t200')
        self.assertEqual([(x.page,x.unit_scale) for x in items],[(3,1000),(5,10000)])

    def test_scope_precedes_source_and_notes(self):
        items=[LineItem('营业收入','1000',1,ExtractMethod.PDFPLUMBER,'text',statement=True),
               LineItem('营业收入','9999',2,ExtractMethod.PDFPLUMBER,'table'),
               LineItem('营业收入','8888',3,ExtractMethod.PDFPLUMBER,'table')]
        result=resolve(items,{1,2},Gaap.CAS,1,'audit',ExtractMethod.PDFPLUMBER)
        self.assertEqual(result.fields['revenue'].value,1000)

    def test_period_basis(self):
        self.assertEqual(period_factor(PeriodType.Q2),2)
        self.assertAlmostEqual(period_factor(PeriodType.Q3),4/3)
        self.assertEqual(period_factor(PeriodType.Q3,'quarter'),4)
        self.assertEqual(period_factor(PeriodType.Q4),1)
        self.assertEqual(period_factor(PeriodType.Q4,'quarter'),4)

    def test_real_pdf_pipeline_with_two_statement_pages(self):
        import tempfile
        from pathlib import Path
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        with tempfile.TemporaryDirectory() as tmp:
            writer=PdfWriter()
            pages=[['Consolidated balance sheet','Amounts in dollars',
                    'Current assets 600','Current liabilities 200',
                    'Total assets 1000','Total liabilities 400','Total equity 600','Retained earnings 200'],
                   ['Consolidated income statement','Amounts in dollars','Revenue 1000','Profit before tax 90','Interest expense 10']]
            for lines in pages:
                page=writer.add_blank_page(width=612,height=792)
                font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
                page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
                stream=DecodedStreamObject()
                lines=lines+['Financial report prepared for regression testing only.']*3
                stream.set_data(('BT /F1 12 Tf 40 750 Td 18 TL '+' '.join('('+line+') Tj T*' for line in lines)+' ET').encode())
                page[NameObject('/Contents')]=writer._add_object(stream)
            filename=Path(tmp)/'report.pdf'
            writer.write(filename)
            result=process(self.profile(),str(filename))
            self.assertAlmostEqual(result['zscore'].z_score,2.3949)
            self.assertEqual(result['extraction'].fields['interest_expense'].value,10)
            self.assertEqual(result['extraction'].fields['ebit'].source_page,1)

    def test_english_current_label_and_chinese_million(self):
        from zscore_pipeline.m2_extract import explicit_unit_scale
        items=table_items([['Item','Current year'],['Current assets','600'],['Interest expense','10']],0,ExtractMethod.PPTX)
        self.assertEqual([x.raw_value for x in items],['600','10'])
        self.assertEqual(explicit_unit_scale('单位：百万元'),1000000)

    def test_consolidated_scope_stops_at_notes(self):
        from zscore_pipeline.m3_locate import consolidated_pages_from_texts
        pages=['合并资产负债表','续表','财务报表附注','母公司利润表','Consolidated income statement']
        self.assertEqual(consolidated_pages_from_texts(pages),{0,1,4})

    def test_real_pptx_table_preserves_columns(self):
        from pptx import Presentation
        from pptx.util import Inches
        from zscore_pipeline.m2_extract_pptx import _extract_pptx_tables
        prs=Presentation()
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        table=slide.shapes.add_table(3,4,Inches(0),Inches(0),Inches(8),Inches(3)).table
        for row,values in zip(table.rows,[['项目','行次','本期','上期'],['资产总计','1','','900000'],['利息费用','2','500','400']]):
            for cell,value in zip(row.cells,values): cell.text=value
        items=_extract_pptx_tables(slide,0,ExtractMethod.PPTX)
        self.assertEqual([(it.label,it.raw_value) for it in items],[('资产总计',''),('利息费用','500')])

    def test_ocr_tab_and_markdown_preserve_empty_current(self):
        for body in ['营业收入\t\t900000', '| 项目 | 附注 | 本期 | 上期 |\n| 营业收入 | 5 | | 900000 |']:
            items=_parse_ocr_text('单位：元\n'+body)
            ext=resolve(items,set(),Gaap.CAS,1,'audit',ExtractMethod.GLM_OCR)
            self.assertNotIn('revenue',ext.fields)

    def test_ocr_retries_unknown_unit_and_invalid_numbers(self):
        for first in ['营业收入\t1000', '单位：元\n营业收入\t1000元']:
            with patch.dict(m2.GLM_OCR,{'verify_attempts':2,'verify_fields':['revenue']}), patch.object(m2,'ocr_pdf_via_glm',side_effect=[first,'单位：千元\n营业收入\t2000']):
                items,_,count,missing=m2._ocr_with_verification('unused')
            self.assertEqual(count,2)
            self.assertEqual(missing,[])
            ext=resolve(items,set(),Gaap.CAS,1,'audit',ExtractMethod.GLM_OCR)
            self.assertEqual(ext.fields['revenue'].value,2000000)

if __name__ == '__main__':
    unittest.main()
