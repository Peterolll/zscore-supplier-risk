# Financial amounts in these fixtures are synthetic.
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from PIL import Image
from zscore_pipeline.pipeline import process
from zscore_pipeline.models import SupplierProfile,ExtractionResult,FieldValue,ExtractMethod,Gaap
from zscore_pipeline.m5_validate import run_gates
from zscore_pipeline.m4_map import resolve
from zscore_pipeline.m2_extract import LineItem
from zscore_pipeline.vision import original_bbox,extract_image,classify_page,read_region


class EvidenceWorkflow(unittest.TestCase):
    def profile(self):
        return SupplierProfile(name='excel-test',industry_class='manufacturing',period_type='annual',gaap='cas')

    def test_excel_two_sides_selects_current_and_preserves_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'report.xlsx';wb=Workbook();ws=wb.active;ws.title='2025资产负债表'
            ws.append(['2025年财务报表','单位：元'])
            ws.append(['资产','行次','期末余额','年初余额','负债和所有者权益','行次','期末余额','年初余额'])
            ws.append(['流动资产合计',14,3000.24,1800.55,'流动负债合计',48,2400.73,1400.63])
            ws.append(['资产总计',34,5400.5,3500.14,'负债合计',60,3200.11,2000.35])
            ws.append(['未分配利润',70,-3100.48,-2600.63,'所有者权益合计',71,2200.39,1499.79])
            ws2=wb.create_sheet('2025利润表');ws2.append(['单位：元']);ws2.append(['项目','本年累计','上年累计'])
            ws2.append(['营业收入',4600.36,2000.70]);ws2.append(['利润总额',-520.39,-650.41]);ws2.append(['利息费用',None,None]);ws2.append(['四、净利润（净亏损以“-”号填列）',-524.85,-650.41]);ws2.append(['所得税费用',4.46,0])
            wb.save(p)
            with patch('zscore_pipeline.ocr_glm._call_glm',side_effect=AssertionError('Excel must not invoke OCR')):
                result=process(self.profile(),str(p))
            f=result['extraction'].fields
            self.assertEqual(f['current_assets'].value,3000.24)
            self.assertEqual(f['current_liabilities'].value,2400.73)
            self.assertEqual(f['net_profit'].value,-524.85)
            self.assertEqual(f['current_assets'].provenance[0]['cell'],'C3')
            self.assertEqual(f['current_assets'].provenance[0]['bbox'],[3,3,3,3])
            self.assertEqual(f['current_assets'].provenance[0]['unit'],'元')
            self.assertIn('2025',f['current_assets'].provenance[0]['period'])
            self.assertNotIn('ebit',f)
            self.assertIsNone(result['zscore'].z_score)
            self.assertIn('G6_净利润勾稽',result['zscore'].gates_passed)
            self.assertIsNotNone(result['extraction'].trace_id)

    def test_excel_formula_no_cache_is_not_zero(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'formula.xlsx';wb=Workbook();ws=wb.active
            ws.append(['单位：元']);ws.append(['项目','金额（元）']);ws.append(['资产总计','=1+2']);wb.save(p)
            result=process(self.profile(),str(p))
            self.assertNotIn('total_assets',result['extraction'].fields)
            self.assertTrue(any('公式无缓存' in x for x in result['extraction'].warnings))

    def test_user_four_column_excel(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'pairs.xlsx';wb=Workbook();ws=wb.active
            ws.append(['项目','金额（元）','项目','金额（元）'])
            ws.append(['流动资产合计',3000.24,'流动负债合计',2400.73]);wb.save(p)
            result=process(self.profile(),str(p))
            self.assertEqual(result['extraction'].fields['current_liabilities'].value,2400.73)

    def test_derived_equity_cannot_pass_identity_check(self):
        f={'total_assets':FieldValue(value=1000),'total_liabilities':FieldValue(value=400),
           'equity_total':FieldValue(value=600,method=ExtractMethod.DERIVED,dependencies=['total_assets','total_liabilities'])}
        passed,failed,notes=run_gates(ExtractionResult(profile_name='x',method='pdfplumber',fields=f))
        self.assertNotIn('G1_复式平衡',passed)
        self.assertTrue(any('未独立验证' in n for n in notes))

    def test_latest_period_excludes_previous_year(self):
        items=[LineItem('资产总计','100',0,ExtractMethod.EXCEL,period='2024期末'),LineItem('资产总计','200',1,ExtractMethod.EXCEL,period='2025期末')]
        self.assertEqual(resolve(items,set(),Gaap.CAS,1,'x',ExtractMethod.EXCEL).fields['total_assets'].value,200)

    def test_direction_retry_and_coordinate_transform(self):
        replies=['{"rotation":0,"kind":"BS","confidence":0.2}', '{"rotation":90,"kind":"BS","layout":"two-column","unit":"元","confidence":0.9}']
        with tempfile.TemporaryDirectory() as tmp,patch('zscore_pipeline.ocr_glm._call_glm',side_effect=replies) as call:
            meta=classify_page(Image.new('RGB',(200,400)),Path(tmp),0)
            self.assertEqual(meta['rotation'],90);self.assertEqual(call.call_count,2)
        self.assertEqual(original_bbox([0,0,1,1],(0,0,.5,1),90),[0,0.5,1,1])
        self.assertIsNone(original_bbox([-1,0,1,1],(0,0,1,1),0))

    def test_split_asset_regions_preserve_header_unit(self):
        meta={'rotation':90,'kind':'BS','layout':'two-column','unit':'元','period':'2025期末','confidence':.9}
        with tempfile.TemporaryDirectory() as tmp,patch('zscore_pipeline.vision.classify_page',return_value=meta),patch('zscore_pipeline.vision.read_region',return_value=[]) as read:
            extract_image(Image.new('RGB',(200,400)),Path(tmp),0,'source.pdf')
            self.assertEqual(read.call_count,3)
            self.assertEqual(read.call_args_list[0].args[0].size,(400,200))
            self.assertEqual(read.call_args_list[0].args[2]['unit'],'元')

    def test_unit_header_retry_is_separate(self):
        rows=[{'label':'营业收入','raw':'123.00','column':'本年累计','bbox':[.1,.1,.5,.2]},
              {'label':'利润总额','raw':'-5.00','column':'本年累计','bbox':[.1,.3,.5,.4]}]
        replies=[json.dumps({'rows':rows,'quality':'clear'}),json.dumps({'unit':'元','period':'2025年度'})]
        with tempfile.TemporaryDirectory() as tmp,patch('zscore_pipeline.ocr_glm._call_glm',side_effect=replies) as call:
            items=read_region(Image.new('RGB',(200,400)),(0,0,1,1),{'rotation':0,'kind':'IS'},1,Path(tmp),'x.pdf')
            self.assertEqual(len(items),2);self.assertEqual(call.call_count,2)
            self.assertIn('计量单位',call.call_args_list[1].args[1])
            self.assertEqual(items[0].raw_value,'123.00');self.assertEqual(items[0].period,'2025年度')

    def test_same_year_date_columns_select_latest(self):
        from zscore_pipeline.evidence import select_column
        self.assertEqual(select_column(['项目','2025年1月1日','2025年12月31日']),2)
        items=[LineItem('资产总计','100',0,ExtractMethod.EXCEL,period='2025-01-01'),LineItem('资产总计','200',0,ExtractMethod.EXCEL,period='2025-12-31')]
        self.assertEqual(resolve(items,set(),Gaap.CAS,1,'x',ExtractMethod.EXCEL).fields['total_assets'].value,200)

    def test_blur_and_missing_fields_retry_is_bounded_and_targeted(self):
        good=[{'label':'营业收入','raw':'123','column':'本年累计','bbox':[.1,.3,.5,.4]},{'label':'利润总额','raw':'-5','column':'本年累计','bbox':[.1,.6,.5,.7]}]
        for quality in ['blurred','clear']:
            first={'unit':'元','quality':quality,'rows':good[:1]}
            last={'unit':'元','quality':'clear','rows':good}
            with tempfile.TemporaryDirectory() as tmp,patch('zscore_pipeline.ocr_glm._call_glm',side_effect=[json.dumps(first),json.dumps(last)]) as call:
                items=read_region(Image.new('RGB',(200,400)),(0,0,1,1),{'rotation':0,'kind':'IS','unit':'元'},1,Path(tmp),'x.pdf')
                self.assertEqual(len(items),2);self.assertEqual(call.call_count,2)
                self.assertIn('local_blur' if quality=='blurred' else 'missing_fields',call.call_args_list[1].args[1])

    def test_year_headers_do_not_conflict_or_replace_blank_current(self):
        rows=[{'label':'营业收入','raw':'10','column':'2024年度','period':'2024年度','bbox':[.1,.1,.3,.2]},
              {'label':'营业收入','raw':'20','column':'2025年度','period':'2025年度','bbox':[.4,.1,.6,.2]},
              {'label':'利润总额','raw':'2','column':'2025年度','bbox':[.4,.3,.6,.4]}]
        with tempfile.TemporaryDirectory() as tmp,patch('zscore_pipeline.ocr_glm._call_glm',return_value=json.dumps({'unit':'元','rows':rows})):
            items=read_region(Image.new('RGB',(200,400)),(0,0,1,1),{'rotation':0,'kind':'IS','period':'2025年度','unit':'元'},0,Path(tmp),'x.pdf')
            revenue=[i for i in items if i.label=='营业收入']
            self.assertEqual(len(revenue),1);self.assertEqual(revenue[0].raw_value,'20')

    def test_excel_cached_identity_formula_is_not_independent(self):
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'cached.xlsx';wb=Workbook();ws=wb.active
            ws.append(['单位：元']);ws.append(['项目','金额（元）']);ws.append(['资产总计',1000]);ws.append(['负债合计',400]);ws.append(['所有者权益合计','=B3-B4']);wb.save(p)
            with ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
            root=ET.fromstring(parts['xl/worksheets/sheet1.xml']);ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            for cell in root.findall('.//s:c',ns):
                if cell.attrib.get('r')=='B5':cell.find('s:v',ns).text='600'
            parts['xl/worksheets/sheet1.xml']=ET.tostring(root)
            with ZipFile(p,'w') as z:
                for name,data in parts.items():z.writestr(name,data)
            result=process(self.profile(),str(p));field=result['extraction'].fields['equity_total']
            self.assertEqual(field.value,600)
            self.assertEqual(field.provenance[0]['formula'],'=B3-B4')
            self.assertNotIn('G1_复式平衡',result['zscore'].gates_passed)

    def test_excel_merged_nonstandard_headers_all_sheets_and_missing_inventory(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'merged.xlsx';wb=Workbook();ws=wb.active;ws.title='资产2025'
            ws.append(['编制单位：测试公司',None,46022,'单位：元']);ws.merge_cells('A1:B1')
            ws.append(['财务科目',None,'本期金额',None,'上期金额']);ws.merge_cells('A2:B2');ws.merge_cells('C2:D2')
            ws.append(['资产总计',None,1000,None,800]);ws.merge_cells('A3:B3');ws.merge_cells('C3:D3')
            ws.append(['负债合计',None,None,None,500]);ws.merge_cells('A4:B4')
            ws.append(['未知财务科目',None,12,None,10]);ws.merge_cells('A5:B5')
            ws2=wb.create_sheet('利润2025');ws2.sheet_state='hidden'
            ws2.append(['单位：元']);ws2.append(['名称','行次','本年累计金额','上期累计金额'])
            ws2.append(['一、营业收入',1,777,666]);ws2.append(['利润总额',2,'=1+2',None]);wb.save(p)
            ext=process(self.profile(),str(p))['extraction']
            self.assertEqual(ext.fields['total_assets'].value,1000)
            self.assertEqual(ext.fields['revenue'].value,777)
            self.assertEqual(ext.fields['total_assets'].provenance[0]['cell'],'C3')
            self.assertNotIn('total_liabilities',ext.fields)
            reasons={(r['sheet'],r['cell']):r['reason'] for r in ext.unextracted}
            self.assertIn('为空',reasons[('资产2025','C4')]);self.assertIn('未映射',reasons[('资产2025','C5')]);self.assertIn('公式无缓存',reasons[('利润2025','C4')])

    def test_excel_indented_headerless_pairs_and_ambiguous_amounts(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'pairs.xlsx';wb=Workbook();ws=wb.active
            ws.append(['单位：元']);ws.append([None,'资产总计',None,1000,None,'负债合计',400])
            ws.append([None,'营业收入',10,20]);wb.save(p)
            ext=process(self.profile(),str(p))['extraction']
            self.assertEqual(ext.fields['total_assets'].value,1000);self.assertEqual(ext.fields['total_liabilities'].value,400)
            self.assertNotIn('revenue',ext.fields)
            self.assertTrue(any('金额列不明确' in r['reason'] for r in ext.unextracted))

    def test_excel_multilevel_year_header_selects_current_without_shifting(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'levels.xlsx';wb=Workbook();ws=wb.active
            ws.append(['单位：元']);ws.append(['项目','2025年',None,'2024年']);ws.merge_cells('A2:A3');ws.merge_cells('B2:C2')
            ws.append([None,'期末余额','年初余额','期末余额'])
            ws.append(['资产总计',1000,800,800]);wb.save(p)
            ext=process(self.profile(),str(p))['extraction']
            self.assertEqual(ext.fields['total_assets'].value,1000)

    def test_all_sheets_include_income_before_ai_and_record_empty_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'workbook.xlsx';wb=Workbook();ws=wb.active;ws.title='资产负债表'
            ws.append(['单位：元']);ws.append(['项目','期末余额']);ws.append(['资产总计',1000])
            ws=wb.create_sheet('利润表');ws.append(['单位：元']);ws.append(['名称','行次','本年累计金额','上期累计金额'])
            ws.append(['营业收入',1,4600.36,20]);ws.append(['财务费用',2,96.8,10]);ws.append(['利润总额',3,-520.39,10])
            wb.create_sheet('附注空页');wb.save(p)
            with patch('zscore_pipeline.ocr_glm._call_glm',side_effect=AssertionError('must not call AI')):
                ext=process(self.profile(),str(p))['extraction']
            self.assertEqual(ext.fields['revenue'].value,4600.36)
            self.assertEqual(ext.fields['revenue'].provenance[0]['sheet'],'利润表')
            self.assertEqual(ext.fields['finance_expenses'].value,96.8)
            self.assertNotIn('ebit',ext.fields)
            self.assertEqual([s['sheet'] for s in ext.sheet_summary],['资产负债表','利润表','附注空页'])
            self.assertEqual(ext.sheet_summary[-1]['candidates'],0)

    def test_ebit_routes_require_explicit_interest_and_preserve_direct_value(self):
        def extract(rows):
            return resolve([LineItem(k,str(v),0,ExtractMethod.EXCEL) for k,v in rows],set(),Gaap.CAS,1,'x',ExtractMethod.EXCEL)
        ext=extract([('净利润',-100),('所得税费用',20),('利息支出',5)])
        self.assertEqual(ext.fields['ebit'].value,-75)
        self.assertEqual(ext.fields['ebit'].dependencies,['net_profit','income_tax','interest_expense'])
        self.assertEqual(len(ext.fields['ebit'].provenance),3)
        ext=extract([('利润总额',-80),('财务费用',5)])
        self.assertNotIn('ebit',ext.fields);self.assertNotIn('interest_expense',ext.fields)
        ext=extract([('息税前利润（EBIT）',-70),('利润总额',-80),('利息支出',5)])
        self.assertEqual(ext.fields['ebit'].value,-70);self.assertEqual(ext.fields['ebit'].dependencies,[])
        ext=extract([('净利润',-100),('利息支出',5)])
        self.assertNotIn('ebit',ext.fields)

if __name__=='__main__':unittest.main()
