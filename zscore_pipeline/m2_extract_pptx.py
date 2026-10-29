"""
m2_extract_pptx.py — PPTX 提取通道

用 python-pptx 读取 PPTX 中的表格、文本框和图片，
输出与 m2_extract.extract_pdf 相同的 (LineItem列表, ExtractMethod, text_len) 结构。

PPTX 有两种常见场景：
1. 电子表格型 PPTX：内含 PPT 原生表格或文本框 → 直接提取
2. 截图型 PPTX：财报截图以图片形式嵌入 → 提取图片走 GLM-4V-Flash OCR

判断逻辑：先提取表格+文本，若不足（无财务关键词）则提取图片走 OCR。
"""
from __future__ import annotations
import re
import tempfile
import os
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import ExtractMethod
from .m2_extract import table_items, attach_page_metadata, finalize_units
from .m2_extract import LineItem, _is_numeric, _extract_text_lines, _parse_ocr_text, detect_unit_scale
from .config import GLM_OCR
from .audit import event


def _safe_iter_shapes(slide):
    """安全遍历 slide 中的 shapes，兼容某些 PPTX 的 rId 属性异常。"""
    result = []
    try:
        for shape in slide.shapes:
            result.append(shape)
    except Exception:
        pass
    return result


def _extract_pptx_tables(slide, slide_no: int, method: ExtractMethod) -> list[LineItem]:
    """从 PPTX slide 的表格中提取 LineItem。"""
    items: list[LineItem] = []
    for shape in _safe_iter_shapes(slide):
        try:
            if not shape.has_table:
                continue
        except Exception:
            continue
        try:
            tbl = shape.table
        except Exception:
            continue
        extracted=table_items([[cell.text for cell in row.cells] for row in tbl.rows], slide_no, method)
        for item in extracted:
            item.bbox=[shape.left,shape.top,shape.left+shape.width,shape.top+shape.height]
            item.coordinate_space='pptx_emu'
            item.cell='表格：'+item.label
        items += extracted
    return items


def _extract_pptx_text(slide, slide_no: int, method: ExtractMethod) -> list[LineItem]:
    """从 PPTX slide 的文本框中提取 LineItem，复用 PDF 文本行解析逻辑。"""
    items: list[LineItem] = []
    for shape in _safe_iter_shapes(slide):
        try:
            if shape.has_table:
                continue  # 表格已单独处理
        except Exception:
            continue
        try:
            has_tf = shape.has_text_frame
        except Exception:
            has_tf = False
        if not has_tf:
            continue
        try:
            text = shape.text_frame.text
        except Exception:
            continue
        if not text:
            continue
        extracted=_extract_text_lines(text,slide_no,method)
        for item in extracted:
            item.bbox=[shape.left,shape.top,shape.left+shape.width,shape.top+shape.height]
            item.coordinate_space='pptx_emu'
        items += extracted
    return items


def _extract_pptx_images(slide, slide_no: int, tmp_dir: str) -> list[str]:
    """提取 PPTX slide 中的图片到临时文件，返回文件路径列表。"""
    paths = []
    for shape in _safe_iter_shapes(slide):
        try:
            st = shape.shape_type
        except Exception:
            continue
        # PICTURE 类型 = 13
        if st != MSO_SHAPE_TYPE.PICTURE:
            continue
        try:
            img = shape.image
            ext = img.ext if hasattr(img, 'ext') else 'png'
            fname = f"slide{slide_no:03d}_shape{id(shape):03d}.{ext}"
            fpath = os.path.join(tmp_dir, fname)
            with open(fpath, 'wb') as f:
                f.write(img.blob)
            paths.append(fpath)
        except Exception:
            # 某些 shape 的 image 属性可能因 rId 问题失败，跳过
            continue
    return paths


# 合并报表关键词
_CONS_KW = ["合并资产负债表", "合并利润表", "合并现金流量表",
            "合併資產負債表", "合併利潤表", "合併現金流量表"]
_PARENT_KW = ["母公司资产负债表", "母公司利润表", "母公司现金流量表"]

# 财务报表小计/总计关键词（用于判断电子文本是否足够）
_FS_KW = [
    "资产总计", "资产合计", "负债合计", "负债总计", "所有者权益合计", "权益总计",
    "資產總計", "負債總計", "權益總額",
    "流动资产合计", "流动负债合计",
    "流動資產合計", "流動負債合計",
    "总计", "總計", "合計",
    "total assets", "total liabilities", "total equity",
]


def _collect_slide_text(slide) -> str:
    """收集 slide 中所有文本（含表格单元格）。"""
    text = ""
    for shape in _safe_iter_shapes(slide):
        try:
            if shape.has_text_frame:
                text += shape.text_frame.text + "\n"
        except Exception:
            pass
        try:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text += (cell.text or "") + " "
        except Exception:
            pass
    return text


def detect_pptx_consolidated_slides(pptx_path: str | Path) -> set[int]:
    """返回"合并报表"相关 slide 页码集合（优先口径）。"""
    from .m3_locate import consolidated_pages_from_texts
    prs = Presentation(str(pptx_path))
    return consolidated_pages_from_texts([_collect_slide_text(slide) for slide in prs.slides])


def detect_pptx_unit_scale(pptx_path: str | Path) -> float:
    """检测 PPTX 报表计量单位 → 换算到元。默认 1（元）。"""
    prs = Presentation(str(pptx_path))
    text = ""
    slides = list(prs.slides)
    for slide in slides:
        text += _collect_slide_text(slide) + " "
    return detect_unit_scale(text)


def extract_pptx(pptx_path: str | Path) -> tuple[list[LineItem], ExtractMethod, int]:
    """提取 PPTX → (LineItem列表, method, text_len)。

    与 extract_pdf 接口一致，供 pipeline 调用。

    策略：
    1. 先提取表格+文本框（电子文本通道）
    2. 若电子文本含足够财务关键词 → 返回
    3. 否则提取图片 → GLM-4V-Flash OCR → 解析（图片型 PPTX 回退）
    """
    pptx_path = str(pptx_path)
    items: list[LineItem] = []
    text_len = 0

    prs = Presentation(pptx_path)
    for i, slide in enumerate(prs.slides):
        text = _collect_slide_text(slide)
        page_items = _extract_pptx_tables(slide, i, ExtractMethod.PPTX)
        page_items += _extract_pptx_text(slide, i, ExtractMethod.PPTX)
        attach_page_metadata(page_items, text)
        for item in page_items:item.source_file=str(Path(pptx_path).resolve())
        items += page_items
        text_len += len(text)

    finalize_units(items)
    # 判断电子文本是否足够
    joined = " ".join(it.label for it in items)
    electronic_enough = text_len > 300 and any(k in joined.lower() for k in _FS_KW)

    if electronic_enough:
        return items, ExtractMethod.PPTX, text_len

    # 电子文本不足 → 提取图片走 GLM OCR
    if not GLM_OCR["enabled"]:
        from .ocr_glm import OCRUnavailable
        raise OCRUnavailable(
            f"图片型 PPTX 需 GLM-4V-Flash OCR，但 GLM_API_KEY 未设置：{pptx_path}"
        )

    from .vision import extract_image
    from PIL import Image
    with tempfile.TemporaryDirectory(prefix='pptx_images_') as tmp:
        directory=Path(tmp)
        ocr_items=[]
        for i,slide in enumerate(prs.slides):
            for image_path in _extract_pptx_images(slide,i,tmp):
                try:
                    with Image.open(image_path) as src:image=src.convert('RGB')
                    ocr_items+=extract_image(image,directory,i,str(Path(pptx_path).resolve()))
                except Exception as exc:event('pptx.image_failed',page=i,message=str(exc)[:300])
        return (ocr_items,ExtractMethod.GLM_OCR,sum(len(it.raw_value) for it in ocr_items)) if ocr_items else (items,ExtractMethod.PPTX,text_len)
