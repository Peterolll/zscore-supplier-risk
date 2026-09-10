"""
output.py — 三层产物

1) JSON：机器可读全字段（含溯源 evidence / source_page / method）。
2) Excel：供应商汇总表（因子 / Z / 风险区 / 闸门）。
3) HTML：复核视图（每样本提取证据 + Z + 闸门结果）。
"""
from __future__ import annotations
import json
from pathlib import Path

from .models import ZScoreResult, ExtractionResult


def _fmt(x: float | None) -> str:
    """格式化因子值；None（缺失因子，如不完整提取或 Z'' 无 X5）显示为 '-'。"""
    return "-" if x is None else f"{x:.4f}"

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def write_json(results: list[dict], path: Path):
    payload = []
    for r in results:
        ext: ExtractionResult = r["extraction"]
        z: ZScoreResult = r["zscore"]
        payload.append({
            "profile": r["profile"].model_dump(),
            "extraction": ext.model_dump(),
            "zscore": z.model_dump(),
            "method": r["method"].value if hasattr(r["method"], "value") else str(r["method"]),
        })
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_excel(results: list[dict], path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Z-Score 汇总"
    headers = ["供应商", "模型", "X1", "X2", "X3", "X4", "X5", "Z-Score",
               "风险区", "年化因子", "闸门通过", "闸门失败"]
    ws.append(headers)
    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
    zone_fill = {
        "safe": PatternFill("solid", fgColor="C6EFCE"),
        "grey": PatternFill("solid", fgColor="FFEB9C"),
        "distress": PatternFill("solid", fgColor="FFC7CE"),
    }
    for r in results:
        z: ZScoreResult = r["zscore"]
        p = r["profile"]
        row = [p.name, z.model,
               round(z.X1, 4) if z.X1 is not None else None,
               round(z.X2, 4) if z.X2 is not None else None,
               round(z.X3, 4) if z.X3 is not None else None,
               round(z.X4, 4) if z.X4 is not None else None,
               round(z.X5, 4) if z.X5 is not None else None,
               round(z.z_score, 4) if z.z_score is not None else None,
               z.risk_zone, z.annualize_factor,
               ",".join(z.gates_passed), ",".join(z.gates_failed)]
        ws.append(row)
        fill = zone_fill.get(z.risk_zone)
        if fill:
            ws.cell(row=ws.max_row, column=9).fill = fill
    # 列宽
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["K"].width = 30
    ws.column_dimensions["L"].width = 24
    wb.save(path)


def write_html(results: list[dict], path: Path):
    rows = []
    for r in results:
        z: ZScoreResult = r["zscore"]
        ext: ExtractionResult = r["extraction"]
        p = r["profile"]
        fld = "".join(
            f"<li><b>{k}</b>: {v.value:,.2f} <span class='ev'>[{v.evidence} · p{v.source_page} · {v.method.value}]</span></li>"
            for k, v in ext.fields.items())
        gates = " ".join(f"<span class='ok'>{g}</span>" for g in z.gates_passed)
        if z.gates_failed:
            gates += " " + " ".join(f"<span class='bad'>{g}</span>" for g in z.gates_failed)
        rows.append(f"""
        <div class="card">
          <h2>{p.name} <small>{z.model} · {z.risk_zone}</small></h2>
          <div class="z">Z = {z.z_score:.4f} {'(流量年化×'+str(z.annualize_factor)+')' if z.annualized else ''}</div>
          <div class="xs">X1={_fmt(z.X1)} X2={_fmt(z.X2)} X3={_fmt(z.X3)} X4={_fmt(z.X4)} X5={_fmt(z.X5)}</div>
          <ul class="flds">{fld}</ul>
          <div class="gates">闸门: {gates}</div>
          {f"<div class='warn'>{' ; '.join(ext.warnings)}</div>" if ext.warnings else ""}
        </div>""")
    html = f"""<!doctype html><html lang="zh"><head><meta charset="utf-8">
    <title>供应商 Z-Score 复核</title>
    <style>
      body{{font-family:-apple-system,Segoe UI,Roboto,sans-serif;background:#f5f6f8;margin:0;padding:24px;color:#1f2329}}
      h1{{font-size:20px}} .card{{background:#fff;border-radius:10px;padding:16px 20px;margin:14px 0;box-shadow:0 1px 3px rgba(0,0,0,.08)}}
      h2{{margin:0 0 8px;font-size:16px}} h2 small{{color:#888;font-weight:400;font-size:12px}}
      .z{{font-size:26px;font-weight:700;color:#2b6cb0}} .xs{{color:#555;font-size:13px;margin:4px 0 10px}}
      .flds{{font-size:13px;line-height:1.7;padding-left:18px}} .ev{{color:#999;font-size:11px}}
      .gates{{font-size:12px;margin-top:8px}} .ok{{color:#1a7f37}} .bad{{color:#cf222e;font-weight:700}}
      .warn{{color:#9a6700;font-size:12px;margin-top:6px}}
    </style></head><body>
    <h1>供应商 Z-Score 自动化复核 · {len(results)} 样本</h1>
    {''.join(rows)}
    </body></html>"""
    path.write_text(html, encoding="utf-8")
