"""Native Excel extraction. Never execute macros/formulas or substitute a prior value."""
from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook
from .audit import event
from .evidence import select_column, is_prior
from .m2_extract import LineItem, explicit_unit_scale
from .models import ExtractMethod

_LABEL_HEADER = re.compile(r'^(?:(?:财务|会计|报表|资产|负债)?项目(?:名称)?|名称|名稱|項目|科目(?:名称)?|会计科目|财务科目|报表项目|資產|资产|负债和所有者权益|item|account(?:name)?)$', re.I)
_META_HEADER = re.compile(r'序号|行次|行号|附注|注释|编号|代码|note|index', re.I)


def extract_excel(filename: str):
    values = load_workbook(filename, data_only=True, read_only=False)
    formulas = load_workbook(filename, data_only=False, read_only=False)
    items, warnings, missed, sheet_summary = [], [], [], []
    def miss(ws, cell, label, reason):
        record=dict(sheet=ws.title,cell=cell.coordinate,label=label,
                    raw_value='' if cell.value is None else str(cell.value),reason=reason)
        if record not in missed: missed.append(record)
    try:
        for page, ws in enumerate(values.worksheets):
            before_items, before_missed = len(items), len(missed)
            raw_ws = formulas[ws.title]
            if ws.max_row * ws.max_column > 500000:
                raise ValueError(f'工作表 {ws.title} 超过 500000 单元格限制')
            # Expand merged headers/labels for navigation, retain the actual anchor for evidence.
            anchors = {}
            for merged in ws.merged_cells.ranges:
                for row in ws.iter_rows(min_row=merged.min_row,max_row=merged.max_row,min_col=merged.min_col,max_col=merged.max_col):
                    for cell in row: anchors[(cell.row,cell.column)]=ws.cell(merged.min_row,merged.min_col)
            def anchor(cell): return anchors.get((cell.row,cell.column),cell)
            def text(cell):
                val=anchor(cell).value
                return str(val).strip() if val is not None else ''
            scale, unit, report_year = None, None, ''
            columns, header_rows = [], []
            for row in ws.iter_rows():
                texts=[text(c) for c in row]
                actual=[str(c.value).strip() if c.value is not None else '' for c in row]
                row_text=' '.join(texts)
                label_cols=[i for i,t in enumerate(texts) if _LABEL_HEADER.fullmatch(re.sub(r'\s+','',t)) and (i==0 or anchor(row[i]).coordinate!=anchor(row[i-1]).coordinate)]
                # Header context only; data cells must not introduce false units/years.
                is_header=bool(label_cols) or not any(isinstance(c.value,(int,float)) and not isinstance(c.value,bool) for c in row)
                if is_header or re.search(r'单位|單位|金额[（(]|金額[（(]', row_text):
                    found=explicit_unit_scale(row_text)
                    if found is not None:
                        scale=found
                        unit={1:'元',1000:'千元',10000:'万元',1000000:'百万元'}.get(scale)
                    years=re.findall(r'(?<!\d)(20\d{2})(?!\d)',row_text+' '+ws.title)
                    if years: report_year=max(years)
                if label_cols:
                    # Adjacent multi-row header: merge its context, including merged parent dates.
                    if header_rows and row[0].row-header_rows[-1][0]<=1:
                        parent=header_rows[-1][1]
                        texts=[' '.join(dict.fromkeys([parent[i],t])).strip() for i,t in enumerate(texts)]
                    header_rows.append((row[0].row,texts))
                    columns=[]
                    for n,start in enumerate(label_cols):
                        stop=label_cols[n+1] if n+1<len(label_cols) else len(texts)
                        segment=texts[start:stop]
                        # One merged amount heading spans blank cells: keep only its anchor.
                        for offset in range(1,len(segment)):
                            cell=row[start+offset]
                            if anchor(cell).column!=cell.column: segment[offset]=''
                        offset=select_column(segment)
                        columns.append((start,start+offset if offset is not None else None,segment[offset] if offset is not None else ''))
                    continue
                # Nested amount headings below a non-merged label header.
                if columns and any(select_column(['项目',t])==1 for t in texts[1:]) and not any(re.search(r'合计|总计|利润|收入|资产|负债',t) for t in actual):
                    new=[]
                    for n,(start,_,_) in enumerate(columns):
                        stop=columns[n+1][0] if n+1<len(columns) else len(texts)
                        parent=header_rows[-1][1] if header_rows else ['']*len(texts)
                        segment=[(parent[i]+' '+texts[i]).strip() for i in range(start,stop)]
                        offset=select_column(segment)
                        new.append((start,start+offset if offset is not None else None,segment[offset] if offset is not None else ''))
                    columns=new
                    continue
                if not columns and re.search(r'公司名称|编制单位|編製單位|单位[：:]|單位[：:]',row_text):
                    continue
                pairs=columns
                if not pairs:
                    # Headerless pairs may be indented, separated by blanks, or repeated left/right.
                    occupied=[i for i,t in enumerate(actual) if t or raw_ws[row[i].coordinate].data_type=='f']
                    pairs=[]
                    for n,i in enumerate(occupied):
                        if not re.search(r'[\u3400-\u9fffA-Za-z]',actual[i]) or raw_ws[row[i].coordinate].data_type=='f': continue
                        stop=next((j for j in occupied[n+1:] if re.search(r'[\u3400-\u9fffA-Za-z]',actual[j]) and raw_ws[row[j].coordinate].data_type!='f'),len(row))
                        numbers=[j for j in occupied if i<j<stop and not isinstance(row[j].value,bool)]
                        if numbers: pairs.append((i,numbers[0] if len(numbers)==1 else None,'金额'))
                handled=set()
                for label_col,value_col,column in pairs:
                    label=texts[label_col]
                    if not label: continue
                    if value_col is None:
                        if any(c.value is not None for c in row[label_col+1:]): miss(ws,anchor(row[label_col]),label,'金额列不明确，未猜测本期列')
                        continue
                    cell=anchor(row[value_col]); raw_cell=raw_ws[cell.coordinate]
                    handled.add(cell.coordinate)
                    if is_prior(column): continue
                    if raw_cell.data_type=='f' and cell.value is None:
                        reason='公式无缓存结果，请用 Excel 重算并保存'
                        warnings.append(f'{ws.title}!{cell.coordinate} {reason}');miss(ws,raw_cell,label,reason)
                        event('excel.formula_without_cache',sheet=ws.title,cell=cell.coordinate)
                        continue
                    if cell.value is None or isinstance(cell.value,bool):
                        if not label.rstrip().endswith(('：',':')) and re.search(r'资产|负债|权益|利润|收入|费用|公积',label): miss(ws,cell,label,'本期单元格为空，不使用上期数替代')
                        continue
                    from .m4_map import parse_number, canonical_field
                    from .synonym import match_component
                    if parse_number(str(cell.value)) is None:
                        miss(ws,cell,label,'金额不是可解析数字');continue
                    if scale is None: miss(ws,cell,label,'单位缺失，金额不写入财务字段')
                    if canonical_field(label) is None and not match_component(label): miss(ws,cell,label,'已读取单元格，但科目未映射到财务指标')
                    items.append(LineItem(label,str(cell.value),page,ExtractMethod.EXCEL,
                        unit_scale=scale,unit_unknown=scale is None,statement=True,
                        formula=str(raw_cell.value) if raw_cell.data_type=='f' else None,
                        source_file=str(Path(filename).resolve()),sheet=ws.title,cell=cell.coordinate,
                        bbox=[cell.column,cell.row,cell.column,cell.row],coordinate_space='excel_cell',
                        column=column,period=(report_year+' '+column).strip(),unit=unit))
                # Any remaining non-period/non-note numeric cell is explicitly accounted for.
                for i,cell in enumerate(row):
                    if cell.value is None or anchor(cell).coordinate in handled: continue
                    if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):
                        heading=header_rows[-1][1][i] if header_rows else ''
                        if _META_HEADER.search(heading) or is_prior(heading): continue
                        if re.fullmatch(r'20\d{2}',str(cell.value)): continue
                        miss(ws,cell,'','未匹配到明确的科目与本期金额列')
            summary=dict(sheet=ws.title,rows=ws.max_row,columns=ws.max_column,hidden=ws.sheet_state!='visible',candidates=len(items)-before_items,unextracted=len(missed)-before_missed)
            sheet_summary.append(summary)
            event('excel.sheet',**summary)
    finally:
        values.close(); formulas.close()
    event('excel.extracted',candidates=len(items),warnings=warnings,unextracted=missed)
    return items, ExtractMethod.EXCEL, sum(len(it.label)+len(it.raw_value) for it in items), warnings, missed, sheet_summary
